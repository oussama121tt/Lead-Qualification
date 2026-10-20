"""Merge rater parts + independent review + settlement into the lead-rating outputs.

Inputs, all in <folder>/output/_parts/:
  R*.json / L*.json   rater arrays (one object per dossier, RATING_PROMPT.md schema)
  review.json         reviewer array: [{"file", "score", "justification"}] for leads first scored >= 6
  final.json          orchestrator settlement: {"<file>": {"final_score": x, "settle_note": "..."}}
                      (only needed where rater and reviewer differ by more than 1 point)

Outputs, in <folder>/output/:
  leads_rated.json    merged array, sorted by final score
  leads_rated.xlsx    one row per agency, schema columns + rater/reviewer/final scores,
                      Status dropdown and Notes, hyperlinks, colour-scaled score, frozen header

    python tools/build_rated_outputs.py --folder "C:/Users/w26be/Desktop/Toumi/Luxembourg_Agency_Research"

Re-rating under a new prompt: --out output_v2 --prev output/leads_rated.json writes to a
separate folder and adds the previous score plus a "biggest moves" section.
"""
import argparse
import json
import sys
from pathlib import Path

SCHEMA = ["file", "agency", "contact", "role", "real_decision_maker", "city", "country", "website", "linkedin",
          "phone", "score", "tier", "why_fits", "why_not_higher", "clients_regulated", "dev_depth", "ai_today",
          "competitor_risk", "build_opportunity", "offer_focus", "linkedin_activity", "best_hook", "angle", "door_opener", "first_step",
          "second_contact", "evidence_notes"]
STATUSES = ["Not started", "Followed", "Request sent", "Connected", "Messaged", "Call booked",
            "Review/demo sent", "Partner", "Paused", "Lost"]


def tier_for(score: float) -> str:
    if score >= 7:
        return "A"
    if score >= 5:
        return "B"
    if score >= 3:
        return "C"
    return "Out"


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, list):
        return "\n".join(f"- {x}" for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return str(v)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--folder", required=True)
    ap.add_argument("--out", default="output", help="output folder name inside --folder")
    ap.add_argument("--prev", default="", help="previous leads_rated.json (relative to --folder) to compare against")
    args = ap.parse_args()
    out = Path(args.folder) / args.out
    prev = {}
    if args.prev:
        prev = {o["file"]: o["score"] for o in
                json.loads((Path(args.folder) / args.prev).read_text(encoding="utf-8"))}
    parts = out / "_parts"

    rows: dict[str, dict] = {}
    for p in sorted(parts.glob("*.json")):
        if p.stem in ("review", "final"):
            continue
        data = json.loads(p.read_text(encoding="utf-8"))
        for obj in data:
            rows[obj["file"]] = obj
    review = {r["file"]: r for r in json.loads((parts / "review.json").read_text(encoding="utf-8"))} \
        if (parts / "review.json").exists() else {}
    final = json.loads((parts / "final.json").read_text(encoding="utf-8")) if (parts / "final.json").exists() else {}

    merged = []
    for f, r in rows.items():
        first = float(r.get("score") or 0)
        rv = review.get(f)
        rev = float(rv["score"]) if rv and rv.get("score") is not None else None
        if f in final:
            score, note = float(final[f]["final_score"]), final[f].get("settle_note", "")
        elif rev is not None and abs(rev - first) > 1:
            raise SystemExit(f"{f}: rater {first} vs reviewer {rev} differ by more than 1; settle it in final.json")
        elif rev is not None:
            score, note = round((first + rev) / 2 * 2) / 2, "rater and reviewer within 1 point; averaged"
        else:
            score, note = first, "below 6 at first pass; not reviewed"
        obj = {k: r.get(k, "") for k in SCHEMA if k in r or k not in ("build_opportunity", "offer_focus")}
        obj.update({"score": score, "tier": tier_for(score), "rater_score": first,
                    "reviewer_score": rev, "reviewer_justification": (rv or {}).get("justification", ""),
                    "settle_note": note})
        if prev:
            obj["prev_score"] = prev.get(f)
        merged.append(obj)
    merged.sort(key=lambda o: (-o["score"], o["file"]))
    (out / "leads_rated.json").write_text(json.dumps(merged, ensure_ascii=False, indent=1), encoding="utf-8")

    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    present = {k for o in merged for k in o}
    cols = (["rank"] + SCHEMA[:10] + ["score", "tier"] + (["prev_score"] if prev else [])
            + ["rater_score", "reviewer_score", "settle_note"]
            + [c for c in SCHEMA[12:] if c in present] + ["reviewer_justification", "Status", "Notes"])
    wb = Workbook()
    ws = wb.active
    ws.title = "leads_rated"
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2A44")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, o in enumerate(merged, 1):
        vals = []
        for c in cols:
            if c == "rank":
                vals.append(i)
            elif c == "Status":
                vals.append("Not started")
            elif c == "Notes":
                vals.append("")
            elif c in ("score", "prev_score", "rater_score", "reviewer_score"):
                vals.append(o.get(c))
            else:
                vals.append(_text(o.get(c)))
        ws.append(vals)
        r = ws.max_row
        for key in ("website", "linkedin"):
            cell = ws.cell(row=r, column=cols.index(key) + 1)
            url = str(cell.value or "").strip()
            if url and not url.startswith("http") and "." in url and " " not in url:
                url = "https://" + url
            if url.startswith("http"):
                cell.hyperlink = url
                cell.font = Font(color="0563C1", underline="single")
    n = ws.max_row
    sc = get_column_letter(cols.index("score") + 1)
    ws.conditional_formatting.add(f"{sc}2:{sc}{n}", ColorScaleRule(
        start_type="num", start_value=1, start_color="F8696B",
        mid_type="num", mid_value=5.5, mid_color="FFEB84",
        end_type="num", end_value=10, end_color="63BE7B"))
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    st = get_column_letter(cols.index("Status") + 1)
    dv.add(f"{st}2:{st}{n}")
    widths = {"rank": 6, "file": 30, "agency": 26, "contact": 20, "role": 22, "real_decision_maker": 26, "city": 14,
              "country": 12, "website": 26, "linkedin": 30, "phone": 16, "score": 7, "tier": 6, "rater_score": 7,
              "reviewer_score": 8, "settle_note": 40, "why_fits": 60, "why_not_higher": 50, "clients_regulated": 45,
              "dev_depth": 40, "ai_today": 40, "competitor_risk": 35, "build_opportunity": 50, "offer_focus": 40, "prev_score": 7, "linkedin_activity": 30, "best_hook": 50,
              "angle": 55, "door_opener": 40, "first_step": 35, "second_contact": 30, "evidence_notes": 45,
              "reviewer_justification": 50, "Status": 16, "Notes": 30}
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 18)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    xlsx = out / "leads_rated.xlsx"
    wb.save(xlsx)
    load_workbook(xlsx)  # must open without errors
    write_summary(out, parts, merged, bool(prev))
    print(f"{len(merged)} leads -> {xlsx}, leads_rated.json, SUMMARY.md (xlsx opens cleanly)")
    print("\nTop 10:")
    for i, o in enumerate(merged[:10], 1):
        was = f" (was {o['prev_score']:g})" if prev and o.get("prev_score") is not None else ""
        print(f"{i:>2}. {o['score']:<4}{was:<10} {o['agency'][:34]:<34} {str(o['angle'])[:110]}")
    return 0


def _one_line(s, n=160) -> str:
    s = " ".join(str(s or "").split())
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"


def write_summary(out: Path, parts: Path, merged: list[dict], compare: bool = False) -> None:
    """SUMMARY.md: ranking, door-opener reuse, shared owners/groups (hand-written
    in _parts/pairs.md), contact corrections, reviewer disagreements."""
    L = ["# Lead rating summary\n"]
    L.append(f"{len(merged)} agencies rated. Tiers: A = 7 and above, B = 5 to 6.5, C = 3 to 4.5, Out = below 3. "
             "Scores are fit as a white-label partner for RuyaTech's offer, not agency quality.\n")
    L.append("## Ranking\n")
    L.append("| # | Agency | Contact | Score | Tier | Angle |")
    L.append("|---|---|---|---|---|---|")
    for i, o in enumerate(merged, 1):
        L.append(f"| {i} | {o['agency']} | {o['contact']} | {o['score']:g} | {o['tier']} | {_one_line(o['angle'], 170)} |")
    L.append("")

    if compare:
        moves = sorted((o for o in merged if o.get("prev_score") is not None and o["score"] != o["prev_score"]),
                       key=lambda o: -abs(o["score"] - o["prev_score"]))
        L.append("## Biggest moves against the previous rating\n")
        if moves:
            L.append("| Agency | Before | Now | Move | Lead with |")
            L.append("|---|---|---|---|---|")
            for o in moves:
                L.append(f"| {o['agency']} | {o['prev_score']:g} | {o['score']:g} | {o['score'] - o['prev_score']:+g} | "
                         f"{_one_line(o.get('offer_focus'), 140)} |")
        else:
            L.append("No score changed.")
        L.append("")

    L.append("## Which door-opener serves which leads\n")
    L.append("Build each asset once and reuse it. Only leads scored 5 or more are listed.\n")
    groups = {"Prototype / AI-app review": [], "AI help-desk demo": [], "Design-to-build handoff": [], "Other": []}
    for o in merged:
        if o["score"] < 5:
            continue
        d = str(o.get("door_opener") or "").lower()
        key = ("Other" if d.startswith(("other", "none"))
               else "Design-to-build handoff" if d.startswith(("design-to-build", "design to build", "handoff"))
               else "Prototype / AI-app review" if d.startswith(("prototype", "ai-app", "app review"))
               else "AI help-desk demo" if d.startswith(("help-desk", "help desk", "ai help"))
               else "Prototype / AI-app review" if "prototype" in d[:60]
               else "AI help-desk demo" if "help-desk" in d[:60]
               else "Design-to-build handoff" if "design-to-build" in d[:60] else "Other")
        groups[key].append(o)
    for k, items in groups.items():
        if not items:
            continue
        L.append(f"**{k}** ({len(items)})\n")
        for o in items:
            L.append(f"- {o['agency']} ({o['score']:g}): {_one_line(o['door_opener'], 190)}")
        L.append("")

    pairs = parts / "pairs.md"
    L.append("## Shared owners, groups and key clients\n")
    L.append(pairs.read_text(encoding="utf-8").strip() if pairs.exists() else "_not written_")
    L.append("")

    L.append("## Listed contact wrong, inactive, or not the decision-maker\n")
    rows = [o for o in merged if str(o.get("real_decision_maker") or "").strip()]
    if rows:
        for o in rows:
            L.append(f"- **{o['agency']}** ({o['score']:g}), sheet contact {o['contact']}: "
                     f"{_one_line(o['real_decision_maker'], 260)}")
    else:
        L.append("None flagged.")
    L.append("")

    L.append("## Reviewer disagreements and how they were settled\n")
    L.append("An independent reviewer, who had not seen the first scores, re-scored every lead first rated 6 or more. "
             "Within 1 point: averaged. More than 1 point apart: settled by hand, reasons below.\n")
    L.append("| Agency | Rater | Reviewer | Final | How |")
    L.append("|---|---|---|---|---|")
    for o in merged:
        if o.get("reviewer_score") is None:
            continue
        L.append(f"| {o['agency']} | {o['rater_score']:g} | {o['reviewer_score']:g} | {o['score']:g} | "
                 f"{_one_line(o['settle_note'], 300)} |")
    L.append("")
    (out / "SUMMARY.md").write_text("\n".join(L), encoding="utf-8")


if __name__ == "__main__":
    sys.exit(main())
