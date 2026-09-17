"""Export the founders worth emailing, with every detail behind the verdict.

One row per non-duplicate lead whose latest verdict is a target segment
(see the profile [segments]) with a valid Sonnet verdict. Buckets:

  A  strong fit          confidence >= 0.8                       -> approve
  B1 founder fits        founder settled, build unknown          -> 10-second check
  B2 other target        rest of the targets / flagged reason    -> read the reason

Columns carry the evidence the model used (career history, deterministic
signals, quotes, hooks with the exact citation, budget, sensitive data) so a
human or another model can re-judge the decision from the sheet alone.

    python tools/export_worth_emailing.py            # exports/worth_emailing.xlsx + .csv
"""
import csv
import json
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)
_env = dict(re.findall(r'^\s*([A-Z0-9_]+)\s*=\s*"?([^"\n]*)"?\s*$', (ROOT / ".env").read_text(encoding="utf-8"), re.M))
for k, v in _env.items():
    os.environ.setdefault(k, v)

import db as dbmod  # noqa: E402
from profile import load_profile  # noqa: E402

TARGET = tuple(sorted(load_profile().target_segments))


def _j(v, default):
    if v is None or v == "":
        return default
    if isinstance(v, (list, dict)):
        return v
    try:
        return json.loads(v)
    except (json.JSONDecodeError, TypeError):
        return default


def _lines(items) -> str:
    return "\n".join(str(i) for i in items if i not in (None, ""))


def _sequence_variant(offer: str, sensitive: bool) -> str:
    if sensitive:
        entry = load_profile().sequences.offers.get(offer)
        if entry is not None and entry.sensitive_id:
            return f"{offer}_sensitive"
    return offer


def bucket(r) -> str:
    c = float(r.get("confidence") or 0)
    if c >= 0.8:
        return "A strong"
    # B1 is the profile-agnostic "founder fits, build unknown" slice: the
    # founder question is settled but the build method is open, so a quick
    # check beats reading the whole reason. (Under ruyatech this also
    # catches settled-technical-founder slices the old segment-named rule
    # left in B2; both buckets are human review either way.)
    if (r["segment"] in TARGET and r.get("founder_profile") not in (None, "", "unknown")
            and r.get("build_evidence") == "unknown"):
        return "B1 founder fits, build unknown"
    return "B2 other target, review"


def main() -> int:
    conn = dbmod.get_connection()
    try:
        rows = [dict(r) for r in conn.execute("""
            SELECT l.id, l.first_name, l.last_name, l.title, l.company_name, l.email, l.website_url,
                   l.linkedin_url, l.apollo_person, l.apollo_org, l.coverage_notes, l.domain_mismatch,
                   l.domain_mismatch_reason, l.session_id,
                   s.segment, s.confidence, s.founder_profile, s.build_evidence, s.recommended_offer,
                   s.company_stage, s.budget_signal, s.budget_evidence, s.budget_blockers,
                   s.sensitive_data_categories, s.data_sensitivity_score, s.built_with_ai_signals,
                   s.technical_signals, s.pain_signals, s.evidence_quotes, s.personalization_hooks,
                   s.disqualify_reason, s.needs_human_review, s.scored_at,
                   t.app_builder_fingerprint, t.site_builder_fingerprint, t.on_builder_subdomain,
                   t.generator_fingerprint, t.ai_authorship_disclosures_found, t.github_repo_url
            FROM leads l
            JOIN lead_scores s ON s.id = (SELECT MAX(id) FROM lead_scores WHERE lead_id = l.id)
            LEFT JOIN lead_technical_signals t ON t.id = (SELECT MAX(id) FROM lead_technical_signals WHERE lead_id = l.id)
            WHERE l.is_duplicate = 0
            ORDER BY l.id
        """).fetchall()]
    finally:
        conn.close()

    out = []
    for r in rows:
        if r["segment"] not in TARGET:
            continue
        if (r.get("disqualify_reason") or "").startswith(("json_parse_failed", "api_error", "no_content_scraped")):
            continue
        person = _j(r.get("apollo_person"), {}) or {}
        org = _j(r.get("apollo_org"), {}) or {}
        history = person.get("employment_history") or []
        career = "\n".join(
            f"{e.get('title') or '?'} @ {e.get('organization') or '?'} ({e.get('start') or '?'} to {e.get('end') or ('now' if e.get('current') else '?')})"
            for e in history[:8] if isinstance(e, dict))
        hooks = _j(r.get("personalization_hooks"), []) or []
        hook_lines = []
        for h in hooks:
            if isinstance(h, dict):
                hook_lines.append(f"{h.get('hook')}  <- \"{(h.get('based_on') or '')[:160]}\"")
            else:
                hook_lines.append(str(h))
        empty_sensitive = load_profile().sensitive.empty_sentinel
        sens = [c for c in (_j(r.get("sensitive_data_categories"), []) or []) if c and c != empty_sensitive]
        signals = []
        for k in ("app_builder_fingerprint", "site_builder_fingerprint", "generator_fingerprint"):
            if r.get(k):
                signals.append(f"{k}={r[k]}")
        if r.get("on_builder_subdomain"):
            signals.append("on_builder_subdomain=yes")
        if r.get("ai_authorship_disclosures_found"):
            signals.append(f"ai_authorship={r['ai_authorship_disclosures_found']}")
        if r.get("github_repo_url"):
            signals.append(f"github={r['github_repo_url']}")
        b = bucket(r)
        out.append({
            "bucket": b,
            "suggested_action": {"A strong": "APPROVE", "B1 founder fits, build unknown": "CHECK 10s then approve/reject",
                                 "B2 other target, review": "READ reason then decide"}[b],
            "lead_id": r["id"],
            "founder": " ".join(filter(None, [r.get("first_name"), r.get("last_name")])),
            "title": r.get("title") or "",
            "company": r.get("company_name") or "",
            "email": r.get("email") or "",
            "website": r.get("website_url") or "",
            "linkedin": r.get("linkedin_url") or "",
            "segment": r["segment"],
            "confidence": r.get("confidence"),
            "recommended_offer": r.get("recommended_offer") or "",
            # Sensitive variant (<offer>_sensitive) when the profile defines
            # a sensitive sequence for the offer and sensitive data is set.
            "sequence_variant": _sequence_variant(r.get("recommended_offer") or "", sens),
            "founder_profile": r.get("founder_profile") or "unknown",
            "build_evidence": r.get("build_evidence") or "unknown",
            "company_stage": r.get("company_stage") or "",
            "needs_human_review": "yes" if r.get("needs_human_review") else "no",
            "review_reason": r.get("disqualify_reason") or "",
            "career_history": career,
            "headline": person.get("headline") or "",
            "company_facts": ", ".join(f"{k}={org.get(k)}" for k in ("employees", "founded_year", "industry") if org.get(k) is not None),
            "deterministic_signals": "; ".join(signals),
            "built_with_ai_signals": _lines(_j(r.get("built_with_ai_signals"), [])),
            "technical_signals": _lines(_j(r.get("technical_signals"), [])),
            "pain_signals": _lines(_j(r.get("pain_signals"), [])),
            "evidence_quotes": _lines(_j(r.get("evidence_quotes"), [])),
            "hooks (hook <- exact citation)": "\n".join(hook_lines),
            "budget_signal": r.get("budget_signal") or "",
            "budget_evidence": _lines(_j(r.get("budget_evidence"), [])),
            "budget_blockers": _lines(_j(r.get("budget_blockers"), [])),
            "sensitive_data": ", ".join(sens),
            "data_sensitivity_score": r.get("data_sensitivity_score"),
            "domain_mismatch": (r.get("domain_mismatch_reason") or "") if r.get("domain_mismatch") else "",
            "coverage_notes": _lines(_j(r.get("coverage_notes"), [])),
            "scored_at": (r.get("scored_at") or "")[:19],
        })

    order = {"A strong": 0, "B1 founder fits, build unknown": 1, "B2 other target, review": 2}
    out.sort(key=lambda x: (order[x["bucket"]], -(x["confidence"] or 0), x["company"]))

    exports = ROOT / "exports"
    exports.mkdir(exist_ok=True)
    cols = list(out[0].keys()) if out else []
    csv_path = exports / "worth_emailing.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(out)

    xlsx_path = exports / "worth_emailing.xlsx"
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = "worth_emailing"
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDDDDD")
        fills = {"A strong": "C8E6C9", "B1 founder fits, build unknown": "FFF3CD", "B2 other target, review": "E3F2FD"}
        for row in out:
            ws.append([row[c] for c in cols])
            fill = PatternFill("solid", fgColor=fills[row["bucket"]])
            for c in ws[ws.max_row][:2]:
                c.fill = fill
        widths = {"bucket": 26, "suggested_action": 28, "founder": 20, "title": 22, "company": 30, "email": 30,
                  "website": 30, "career_history": 60, "evidence_quotes": 70, "hooks (hook <- exact citation)": 80,
                  "review_reason": 50, "coverage_notes": 50, "built_with_ai_signals": 40, "technical_signals": 40,
                  "pain_signals": 40, "budget_evidence": 40, "budget_blockers": 40, "deterministic_signals": 40}
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 16)
        for r_ in ws.iter_rows(min_row=2):
            for c in r_:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        # Summary sheet
        s2 = wb.create_sheet("summary")
        from collections import Counter
        s2.append(["bucket", "count"])
        for k, n in Counter(x["bucket"] for x in out).most_common():
            s2.append([k, n])
        s2.append([])
        s2.append(["offer", "count"])
        for k, n in Counter(x["recommended_offer"] for x in out).most_common():
            s2.append([k, n])
        s2.append([])
        s2.append(["sequence_variant", "count"])
        for k, n in Counter(x["sequence_variant"] for x in out).most_common():
            s2.append([k, n])
        s2.append([])
        s2.append(["How to read", ""])
        s2.append(["founder_profile", "from Apollo career history: non_technical / semi_technical / technical / unknown"])
        s2.append(["build_evidence", "from site fingerprints + text: ai_built / hand_built / unknown"])
        s2.append(["confidence", "< 0.7 is always routed to human review; 0.5-0.7 = founder settled, build open"])
        s2.append(["hooks", "each hook is followed by the exact source citation it rests on (<- \"...\")"])
        s2.append(["review_reason", "why a human must look: derived segment, dropped citations, domain mismatch, budget blocker"])
        wb.save(xlsx_path)
    except ImportError:
        xlsx_path = None

    from collections import Counter
    print(f"rows: {len(out)} | {dict(Counter(x['bucket'] for x in out))}")
    print("csv :", csv_path)
    print("xlsx:", xlsx_path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
