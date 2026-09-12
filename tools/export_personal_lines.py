"""Review sheet for the generated Personal Lines — the one sentence that opens
the Apollo sequence for each founder.

One row per lead that has a line. The job when reading it: mark KEEP or CUT in
the verdict column. A CUT lead still gets the sequence, just without a
personalised opener, so cutting costs nothing but the personalisation.

    python tools/export_personal_lines.py
"""
import csv
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import db as dbmod


def main() -> int:
    conn = dbmod.get_connection()
    try:
        rows = [dict(r) for r in conn.execute("""
            SELECT l.id, l.first_name, l.last_name, l.title, l.company_name, l.email, l.website_url,
                   l.personal_line, l.personal_line_based_on, l.personal_line_status,
                   s.segment, s.confidence, s.recommended_offer, s.sensitive_data_categories,
                   s.founder_profile, s.build_evidence
            FROM leads l
            JOIN lead_scores s ON s.id = (SELECT MAX(id) FROM lead_scores WHERE lead_id = l.id)
            WHERE l.is_duplicate = 0 AND l.personal_line_status IS NOT NULL
            ORDER BY (l.personal_line_status = 'ok') DESC, COALESCE(s.confidence, 0) DESC, l.id
        """).fetchall()]
    finally:
        conn.close()

    out = []
    for r in rows:
        sens = r.get("sensitive_data_categories")
        try:
            sens = [c for c in (json.loads(sens) if isinstance(sens, str) else (sens or [])) if c and c != "none"]
        except (json.JSONDecodeError, TypeError):
            sens = []
        out.append({
            "your_verdict (KEEP / CUT)": "",
            "why (optional)": "",
            "personal_line": r.get("personal_line") or "",
            "founder": " ".join(filter(None, [r.get("first_name"), r.get("last_name")])),
            "company": r.get("company_name") or "",
            "title": r.get("title") or "",
            "website": r.get("website_url") or "",
            "based_on (the quote it rests on)": r.get("personal_line_based_on") or "",
            "status": r.get("personal_line_status") or "",
            "segment": r.get("segment") or "",
            "confidence": r.get("confidence"),
            "offer": r.get("recommended_offer") or "",
            "sensitive_data": ", ".join(sens),
            "founder_profile": r.get("founder_profile") or "",
            "build_evidence": r.get("build_evidence") or "",
            "email": r.get("email") or "",
            "lead_id": r["id"],
        })

    exports = ROOT / "exports"
    exports.mkdir(exist_ok=True)
    cols = list(out[0].keys()) if out else []
    with open(exports / "personal_lines.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(out)

    xlsx = exports / "personal_lines.xlsx"
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        wb = Workbook()
        ws = wb.active
        ws.title = "personal_lines"
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDDDDD")
        for row in out:
            ws.append([row[c] for c in cols])
            if row["status"] != "ok":
                for c in ws[ws.max_row]:
                    c.fill = PatternFill("solid", fgColor="F0F0F0")
        dv = DataValidation(type="list", formula1='"KEEP,CUT"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"A2:A{ws.max_row}")
        widths = {"your_verdict (KEEP / CUT)": 22, "why (optional)": 30, "personal_line": 95,
                  "founder": 20, "company": 30, "title": 22, "website": 30,
                  "based_on (the quote it rests on)": 60, "status": 14, "segment": 20,
                  "sensitive_data": 24, "email": 30}
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 15)
        for r_ in ws.iter_rows(min_row=2):
            for c in r_:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        s2 = wb.create_sheet("how to read")
        for line in [
            ("What this is", "The one sentence that follows 'Hi <name>,' in email 1 of your Apollo sequence."),
            ("", "Everything else in the sequence is your existing copy, unchanged."),
            ("Your job", "Put KEEP or CUT in column A. Skim the line, glance at the quote it rests on."),
            ("KEEP", "The line is true, specific, and sounds like something you would write."),
            ("CUT", "Generic, wrong, presumptuous, or it reads like advice or a compliment."),
            ("Cost of CUT", "None. That founder still gets the sequence, just with no personalised opener."),
            ("", ""),
            ("The style it copies", "Your own sent lines, e.g.:"),
            ("", "Ove is built for girls going through puberty, so the data belongs to minors and UK rules treat that more strictly than almost anything else."),
            ("", "You're running a full real-estate practice and building Paced at the same time, which usually means the app gets your evenings and the plumbing underneath gets whatever's left."),
            ("", ""),
            ("Guards already applied", "No compliments, no pitch, no advice, no questions, no exclamation marks."),
            ("", "Every line must quote real evidence from their site or profile, or it is discarded."),
            ("Greyed rows", "status is not 'ok' — no line will be sent for that founder."),
        ]:
            s2.append(list(line))
        s2.column_dimensions["A"].width = 24
        s2.column_dimensions["B"].width = 120
        for r_ in s2.iter_rows():
            for c in r_:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        wb.save(xlsx)
    except ImportError:
        xlsx = None

    from collections import Counter
    print(f"rows: {len(out)} | {dict(Counter(x['status'] for x in out))}")
    print("csv :", exports / "personal_lines.csv")
    print("xlsx:", xlsx)
    return 0


if __name__ == "__main__":
    sys.exit(main())
